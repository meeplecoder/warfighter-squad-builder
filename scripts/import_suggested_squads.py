#!/usr/bin/env python3
"""
Import suggested squads from 'Warfighter missions squads.xlsx' (Missions tab).
Run from warfighter-squad-builder/ directory:
    python scripts/import_suggested_squads.py [--xlsx path/to/file.xlsx] [--dry-run]

Re-running replaces existing admin squads with the same name.
"""
import os
import sys
import re
import argparse

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(os.path.dirname(os.path.dirname(__file__)), '.env'))
except ImportError:
    pass

try:
    import openpyxl
except ImportError:
    print('ERROR: openpyxl not installed. Run: pip install openpyxl')
    sys.exit(1)

from werkzeug.security import generate_password_hash
from app import create_app, db
from app.models import Card, CardAssociation, Squad, SquadSoldier, SquadGear, SquadObjective, User
from app.models import now_utc

DEFAULT_NATION = 'US'
ADMIN_USERNAME = 'admin'
ADMIN_EMAIL = 'admin@warfighter.local'


# ---------------------------------------------------------------------------
# Soldier + gear parsing
# ---------------------------------------------------------------------------

_TYPE_PREFIX = re.compile(r'^(PS|NPS|SqS|Sqd)\s*:?\s*', re.IGNORECASE)
_SUBTYPE_MAP = {'PS': 'Player', 'NPS': 'NPS', 'SQS': 'Squad', 'SQD': 'Squad'}


def _clean_name(tok):
    tok = re.sub(r'\(\d+\)', '', tok)       # strip RP cost "(13)"
    tok = re.sub(r'\{[^}]*\}', '', tok)     # strip {role - ...}
    tok = tok.strip(' .,;')
    return tok.strip()


def parse_gear_items(gear_text):
    """Parse comma-separated gear text into [{'name': str, 'qty': int}]."""
    items = []
    for part in gear_text.split(','):
        part = part.strip().rstrip('.')
        if not part:
            continue
        # Remove role hints (closed or unclosed brace)
        part = re.sub(r'\{[^}]*\}', '', part).strip()
        part = re.sub(r'\{[^}]*$', '', part).strip()
        if not part:
            continue
        # Strip trailing RP total marker: ". = RP 35" attached to last item on line
        part = re.sub(r'\.\s*=\s*RP\s*\d+.*$', '', part).strip().rstrip('.')
        # Skip standalone RP totals like "= RP 15" or "RP 15"
        if re.match(r'^=?\s*RP\s*\d+', part, re.IGNORECASE):
            continue
        # Parse leading quantity: "2 Canteens", "3 Mk2 grenades"
        qty = 1
        qm = re.match(r'^(\d+)\s+(.+)', part)
        if qm:
            qty = int(qm.group(1))
            part = qm.group(2).strip()
        # Strip trailing RP cost suffix: "M1 Rifle (2)" or "Canteen (1)"
        part = re.sub(r'\s*\(\d+\)$', '', part).strip()
        if not part or len(part) < 2:
            continue
        items.append({'name': part, 'qty': qty})
    return items


# Name aliases: spreadsheet names → canonical DB card names (lowercase)
_GEAR_ALIASES = {
    'm9 rifle grenade': 'm9 at rifle grenade',
    'm9 at rifle grenade launcher': 'm9 at rifle grenade',
}


def parse_soldier_entries(text):
    """
    Return list of dicts: {'name': str, 'subtype': str, 'gear': [{'name': str, 'qty': int}]}
    """
    if not text:
        return []
    text = str(text)
    results = []

    # Merge continuation lines into the line with the type prefix
    raw_lines = text.split('\n')
    merged = []
    cur = ''
    for line in raw_lines:
        line = line.strip()
        if not line:
            continue
        has_prefix = _TYPE_PREFIX.match(line)
        is_name_first = re.match(r'^[A-Z][a-z]+\s*,\s*(PS|NPS|SqS)', line, re.IGNORECASE)
        if has_prefix or is_name_first:
            if cur:
                merged.append(cur)
            cur = line
        else:
            cur = (cur + ' ' + line).strip()
    if cur:
        merged.append(cur)

    for line in merged:
        # Format A: "PS : Name, ..." / "NPS : Name1 , ... Name2 , ..."
        m = _TYPE_PREFIX.match(line)
        if m:
            type_key = m.group(1).upper()
            subtype = _SUBTYPE_MAP.get(type_key, 'Player')
            rest = line[m.end():].strip()

            # Split multiple soldiers on one line at ". " followed by capital letter
            segments = re.split(r'\.\s+(?=[A-Z])', rest)
            for seg in segments:
                seg = seg.strip().rstrip('.')
                if not seg or seg.lower().startswith('none') or seg.startswith('='):
                    continue
                # Raw name token: up to first whitespace, comma, paren, brace, or period
                raw_name = re.split(r'[\s,({.]', seg)[0]
                name_tok = _clean_name(raw_name)
                if not name_tok or len(name_tok) <= 1:
                    continue

                # Gear: everything after the raw name token + optional RP cost "(N)"
                after_name = seg[len(raw_name):].strip()
                after_name = re.sub(r'^\s*\(\d+\)', '', after_name).strip()
                gear_text = after_name.lstrip(' ,').strip()
                gear_items = parse_gear_items(gear_text)

                results.append({'name': name_tok, 'subtype': subtype, 'gear': gear_items})
            continue

        # Format B: "Name, PS, gear..." (name comes first)
        m2 = re.match(r'^(\w+)\s*,\s*(PS|NPS|SqS|Squad)\b\s*,?\s*(.*)', line, re.IGNORECASE)
        if m2:
            name_tok = _clean_name(m2.group(1))
            type_key = m2.group(2).upper()
            subtype = _SUBTYPE_MAP.get(type_key, 'Player')
            gear_text = m2.group(3).strip()
            if name_tok and len(name_tok) > 1:
                gear_items = parse_gear_items(gear_text)
                results.append({'name': name_tok, 'subtype': subtype, 'gear': gear_items})

    return results


# ---------------------------------------------------------------------------
# Objective name parsing
# ---------------------------------------------------------------------------

_OBJ_HEADER = re.compile(
    r'^(?:EASY|HARD|SPECIAL|MEDIUM)\s+OBJECTIVES?:\s*',
    re.IGNORECASE
)


def _extract_names_from_fragment(fragment, names):
    """Parse comma-separated 'Name* , Name**' tokens into names list."""
    for part in re.split(r',', fragment):
        part = part.strip()
        m = re.match(r'^([^-–*\n]+?)(?:\s*[-–*]|$)', part)
        if m:
            name = m.group(1).strip()
            name = re.sub(r'[\s*\d]+$', '', name).strip(' ,.*:')
            if name and len(name) > 2:
                names.append(name)


def parse_objective_names(text):
    """Extract bare objective card names from objectives column text."""
    if not text:
        return []
    raw = str(text).strip()
    if raw.lower() in ('any', 'none', ''):
        return []
    names = []
    for line in raw.split('\n'):
        line = line.strip()
        if not line:
            continue
        # Lines like "Hard Objectives: Tank Hunting*, Rescue**, ..."
        hm = _OBJ_HEADER.match(line)
        if hm:
            rest = line[hm.end():]
            _extract_names_from_fragment(rest, names)
            continue
        # Plain "Break Out - description" or just "Bunker Clearing"
        _extract_names_from_fragment(line, names)
    return names


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--xlsx', help='Path to Excel file')
    parser.add_argument('--dry-run', action='store_true', help='Parse without writing to DB')
    args = parser.parse_args()

    # Locate xlsx
    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    xlsx_path = args.xlsx or os.path.normpath(os.path.join(base, '..', 'Warfighter missions squads.xlsx'))
    if not os.path.exists(xlsx_path):
        print(f'ERROR: file not found: {xlsx_path}')
        sys.exit(1)

    print(f'Loading: {xlsx_path}')
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.worksheets[0]
    print(f'Sheet: {ws.title!r} | rows: {ws.max_row}')

    app = create_app()
    with app.app_context():
        # Ensure admin user
        admin = User.query.filter_by(username=ADMIN_USERNAME).first()
        if not admin:
            admin = User(
                username=ADMIN_USERNAME,
                email=ADMIN_EMAIL,
                password_hash=generate_password_hash('changeme'),
                full_name='Warfighter Admin',
                email_verified_at=now_utc(),
            )
            if not args.dry_run:
                db.session.add(admin)
                db.session.flush()
            print(f'Admin user created: {ADMIN_USERNAME}')
        else:
            print(f'Admin user: {ADMIN_USERNAME} (id={admin.id})')

        # Build soldier lookup: lower_name → [Card]
        soldier_by_name = {}
        for c in Card.query.filter_by(card_type='Soldier').all():
            soldier_by_name.setdefault(c.name.lower(), []).append(c)

        # Build gear lookup: lower_name → Card, prefer Core Europe
        GEAR_TYPES = ['Weapon', 'Equipment', 'Skill', 'Service Record']
        gear_by_name = {}
        for c in Card.query.filter(Card.card_type.in_(GEAR_TYPES)).all():
            key = c.name.lower()
            existing = gear_by_name.get(key)
            if existing is None or c.module == 'Core Europe':
                gear_by_name[key] = c

        # Build objective lookup: lower_name → Card
        obj_by_name = {}
        for c in Card.query.filter_by(card_type='Objective').all():
            key = c.name.lower()
            if key not in obj_by_name:
                obj_by_name[key] = c

        # Build mission lookup: lower first-line → Card
        mission_by_name = {}
        for c in Card.query.filter_by(card_type='Mission').all():
            key = c.name.lower().split('\n')[0].strip()
            if key not in mission_by_name:
                mission_by_name[key] = c

        def find_soldier(name, subtype_hint):
            candidates = soldier_by_name.get(name.lower(), [])
            if not candidates:
                return None
            typed = [c for c in candidates if c.card_subtype == subtype_hint] if subtype_hint else candidates
            pool = typed or candidates
            us = [c for c in pool if c.nation == 'US']
            return (us or pool)[0]

        def find_gear_card(name):
            key = name.lower().strip()
            # Apply alias map first
            key = _GEAR_ALIASES.get(key, key)
            c = gear_by_name.get(key)
            if c:
                return c
            # Try stripping trailing 's' for plurals: "canteens" → "canteen"
            if key.endswith('s') and len(key) > 2:
                c = gear_by_name.get(key[:-1])
                if c:
                    return c
            return None

        def find_mission(raw):
            first = str(raw).split('\n')[0].strip().lower()
            return mission_by_name.get(first)

        def find_objective(name):
            return obj_by_name.get(name.strip().lower())

        inserted = 0
        replaced = 0
        skip_no_mission = 0
        skip_no_soldiers = 0
        total_gear_miss = 0

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            _packs, mission_raw, rp, _lo, _nsol, soldiers_text, objectives_text, _special, _source = row

            if not soldiers_text or not str(soldiers_text).strip():
                continue
            if not mission_raw:
                continue

            mission_card = find_mission(mission_raw)
            if not mission_card:
                print(f'  Row {row_num}: MISS mission {str(mission_raw)[:50]!r}')
                skip_no_mission += 1
                continue

            # Parse + match soldiers (with gear)
            entries = parse_soldier_entries(str(soldiers_text))
            seen_ids = set()
            matched_entries = []   # [{'card': Card, 'gear': [...]}]
            for entry in entries:
                card = find_soldier(entry['name'], entry['subtype'])
                if card and card.id not in seen_ids:
                    matched_entries.append({'card': card, 'gear': entry['gear']})
                    seen_ids.add(card.id)

            if not matched_entries:
                soldier_names = [(e['name'], e['subtype']) for e in entries]
                print(f'  Row {row_num}: MISS soldiers for {mission_card.name!r} | parsed: {soldier_names}')
                skip_no_soldiers += 1
                continue

            # Parse + match objectives
            obj_names = parse_objective_names(objectives_text)
            seen_obj_ids = set()
            matched_objectives = []
            for oname in obj_names:
                oc = find_objective(oname)
                if oc and oc.id not in seen_obj_ids:
                    matched_objectives.append(oc)
                    seen_obj_ids.add(oc.id)

            # Nation: majority of soldier nations
            nations = [e['card'].nation for e in matched_entries if e['card'].nation]
            nation = max(set(nations), key=nations.count) if nations else DEFAULT_NATION

            mission_display = str(mission_raw).split('\n')[0].strip()
            squad_name = f'Squad for {mission_display}'

            soldier_names_str = ', '.join(e['card'].name for e in matched_entries)
            obj_names_str = ', '.join(c.name for c in matched_objectives) or '—'
            print(f'  Row {row_num}: {squad_name!r} | {soldier_names_str} | obj: {obj_names_str}')

            # Log gear
            for entry in matched_entries:
                for g in entry['gear']:
                    gc = find_gear_card(g['name'])
                    if gc:
                        print(f'    gear: {g["qty"]}x {gc.name} ({gc.card_type})')
                    else:
                        print(f'    GEAR MISS: {g["name"]!r}')
                        total_gear_miss += 1

            if not args.dry_run:
                # Delete existing squad by this admin with same name
                existing = Squad.query.filter_by(user_id=admin.id, name=squad_name).first()
                if existing:
                    db.session.delete(existing)
                    db.session.flush()
                    replaced += 1

                squad = Squad(
                    user_id=admin.id,
                    name=squad_name,
                    mission_card_id=mission_card.id,
                    nation=nation,
                    notes='',
                    rp_total=int(rp) if rp else 0,
                    is_valid=1,
                    status='published',
                )
                db.session.add(squad)
                db.session.flush()

                for j, entry in enumerate(matched_entries):
                    ss = SquadSoldier(squad_id=squad.id, card_id=entry['card'].id, sort_order=j)
                    db.session.add(ss)
                    db.session.flush()
                    # Pre-printed card names for this soldier (skip to avoid duplication)
                    preprinted_names = {
                        a.child_card.name.lower() for a in
                        CardAssociation.query.filter_by(parent_card_id=entry['card'].id).all()
                    }
                    for g in entry['gear']:
                        gc = find_gear_card(g['name'])
                        if gc and gc.name.lower() not in preprinted_names:
                            db.session.add(SquadGear(
                                squad_soldier_id=ss.id,
                                card_id=gc.id,
                                quantity=g['qty'],
                            ))

                for j, oc in enumerate(matched_objectives):
                    db.session.add(SquadObjective(squad_id=squad.id, card_id=oc.id, sort_order=j))

                db.session.commit()

            inserted += 1

        print(f'\n{"DRY RUN — " if args.dry_run else ""}Done.')
        print(f'  Created:              {inserted}')
        print(f'  Replaced existing:    {replaced}')
        print(f'  Skipped (no mission): {skip_no_mission}')
        print(f'  Skipped (no soldier): {skip_no_soldiers}')
        print(f'  Gear misses:          {total_gear_miss}')


if __name__ == '__main__':
    main()
